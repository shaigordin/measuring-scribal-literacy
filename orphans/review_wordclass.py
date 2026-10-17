"""Review, correct and audit the word-class annotation (qtm/wordclass.py on TLHdig's analyses).

The classes are set by rules from TLHdig's morphological analyses, so they inherit both the
editors' choices and the rules' mistakes. Each token carries the rule that set it and a
confidence level (high / medium / low, see qtm.wordclass.RULES). Three steps keep the errors in
view:

  sheet   labels/wordclass_review.csv: one row per word form and class, low and medium confidence
          and forms with more than one class first, then by frequency. Fill `corrected_class`
          (one of the classes below) for wrong rows; leave the rest empty.
  apply   copies the filled rows into labels/wordclass_corrections.csv, which qtm.tlhdig applies
          when it builds the token table (`bash reproduce.sh` reruns everything).
  audit   labels/wordclass_audit.csv: a fixed random sample of tokens per confidence level, with
          the line they stand in. Mark `verdict` y (right) or n (wrong, and give `correct_class`).
  form    labels/wordclass_audit.xlsx: the audit sample as a spreadsheet to fill (rows shuffled,
          y/n and class drop-downs, the confidence level hidden), with instructions.
  collect FILE.xlsx: copies the marks from a filled form back into labels/wordclass_audit.csv.
  compare OLD NEW [OUT]: how a rule change moved the classes between two token tables built from the
          same release (row by row) -> results/wordclass_rule_changes.csv, with examples.
  examples  results/wordclass_examples.csv: for every rule and class, the most frequent word
          forms with a transliteration, the TLHdig analysis the rule read, and a manuscript.
  score   accuracy per confidence level with a 95% Wilson interval, and the corpus-wide estimate
          weighted by the share of tokens at each level -> results/wordclass_audit.json.
  rescore the same marks against the current token table, after rules were revised on them (an
          in-sample upper bound) -> results/wordclass_audit_after_fixes.json.

Usage: python review_wordclass.py {sheet,apply,audit,form,collect,score,rescore,compare,examples} [--tokens data/tokens.csv]
"""
import argparse
import json
import math
from pathlib import Path

import numpy as np
import pandas as pd

from qtm.groups import ABSTRACT, GROUPSETS, HISTORICAL, RESIDUE
from qtm.wordclass import CLASSES

HERE = Path(__file__).parent
SHEET = HERE / "labels" / "wordclass_review.csv"
CORR = HERE / "labels" / "wordclass_corrections.csv"
AUDIT = HERE / "labels" / "wordclass_audit.csv"
FORM = HERE / "labels" / "wordclass_audit.xlsx"
SCORE = HERE / "results" / "wordclass_audit.json"
VALID = set(CLASSES) | {"other", "ambiguous", "unknown"}
LEVELS = ("high", "medium", "low", "manual")
PER_LEVEL = 60
SEED = 0


def tokens(path):
    """Legible words of the texts the experiments use (reference traditions, residue, orphan pools)."""
    t = pd.read_csv(path, low_memory=False, dtype={"cth_number": str})
    used = {str(c) for g in (HISTORICAL, ABSTRACT, *GROUPSETS.values(), RESIDUE) for cs in g.values() for c in cs} | {"470"}
    return t[t.word.notna() & (t.word != "") & t.cth_number.isin(used)]


def sheet(t, top):
    g = t.groupby(["word", "wclass"]).agg(
        translit=("translit", lambda s: s.value_counts().index[0]), rule=("wc_rule", lambda s: s.value_counts().index[0]),
        confidence=("wc_conf", lambda s: s.value_counts().index[0]), tokens=("txtid", "size"),
        texts=("txtid", "nunique")).reset_index()
    n_cls = t.groupby("word").wclass.nunique()
    g["classes_of_this_form"] = g.word.map(n_cls)
    rank = {"low": 0, "medium": 1, "high": 2, "manual": 3}
    g = g.assign(_r=g.confidence.map(rank).fillna(0) - (g.classes_of_this_form > 1)).sort_values(
        ["_r", "tokens"], ascending=[True, False]).drop(columns="_r")
    unsure = g.confidence.isin(["low", "medium"])
    cover = {n: round(float(g.tokens.iloc[:n][unsure.iloc[:n]].sum() / g.tokens[unsure].sum()), 3) for n in (100, 500, top)}
    g = g.iloc[:top].copy()
    g["corrected_class"], g["note"] = "", ""
    if SHEET.exists():  # keep what was already filled in
        old = pd.read_csv(SHEET, dtype=str).fillna("")
        old = old[old.corrected_class != ""].set_index(["word", "wclass"])
        key = list(zip(g.word, g.wclass))
        g["corrected_class"] = [old.corrected_class.get(k, "") for k in key]
        g["note"] = [old.note.get(k, "") for k in key]
    g.to_csv(SHEET, index=False)
    print(f"{SHEET}: {len(g)} rows; share of the low/medium-confidence tokens covered by the first n rows: {cover}")


def apply():
    s = pd.read_csv(SHEET, dtype=str).fillna("")
    s = s[s.corrected_class != ""]
    bad = set(s.corrected_class) - VALID
    if bad:
        raise SystemExit(f"unknown classes {bad}; use one of {sorted(VALID)}")
    cor = pd.read_csv(CORR, dtype=str).fillna("") if CORR.exists() else pd.DataFrame(
        columns=["word", "translit", "wclass", "note", "reviewer"])
    new = pd.DataFrame({"word": s.word, "translit": "", "wclass": s.corrected_class, "note": s.note, "reviewer": ""})
    out = pd.concat([cor, new]).drop_duplicates(["word", "translit"], keep="last")
    out.to_csv(CORR, index=False)
    print(f"{CORR}: {len(out)} corrections ({len(new)} from the review sheet); rerun reproduce.sh")


def audit(t):
    rng = np.random.default_rng(SEED)
    rows = []
    for lev in ("high", "medium", "low"):
        g = t[t.wc_conf == lev]
        rows.append(g.iloc[np.sort(rng.choice(len(g), min(PER_LEVEL, len(g)), replace=False))])
    a = pd.concat(rows)
    line = t.groupby(["txtid", "lnr"]).translit.agg(lambda s: " ".join(map(str, s)))
    a = a.assign(line=[line.get((x, l), "") for x, l in zip(a.txtid, a.lnr)], verdict="", correct_class="")
    a = a[["txtid", "lnr", "line", "word", "translit", "wclass", "wc_rule", "wc_conf", "verdict", "correct_class"]]
    if AUDIT.exists():
        raise SystemExit(f"{AUDIT} exists; delete it to draw a new sample")
    a.to_csv(AUDIT, index=False)
    print(f"{AUDIT}: {len(a)} tokens ({PER_LEVEL} per confidence level) to mark y/n")


CLASS_HELP = [
    ("function", "function word: connective, particle, pronoun, adverb, preverb, postposition, preposition (nu, -ma, kuit, arḫa, A-NA, ina)"),
    ("verb", "finite or non-finite verb form, participles included (memai, dāi, ḫalziya)"),
    ("realia", "noun, adjective or other nominal form: objects, materials, places, abstract nouns (NINDA.GUR₄.RA, GEŠTIN, ḫaššan)"),
    ("agent", "a person or a group acting: personal names, titles, professions ({MUNUS}ŠU.GI, LUGAL, {m}Kupanta-{D}KAL)"),
    ("deity", "a divine name or the word for god ({D}UTU-uš, DINGIR.MEŠ)"),
    ("foreign", "a word in Hurrian, Hattian, Luwian or Palaic (not a logogram in such a line)"),
    ("other", "a number or a word that fits none of the above"),
    ("unknown", "too broken to tell"),
]
FORM_COLS = ["id", "manuscript", "line no.", "line (TLHdig transliteration)", "word as written",
             "proposed category", "correct? (y/n)", "correct category (if n)", "note"]


def form(t):
    """The audit sample as an .xlsx form: shuffled so confidence levels are not grouped."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    a = pd.read_csv(AUDIT, dtype=str).fillna("").reset_index().rename(columns={"index": "id"})
    a = a.sample(frac=1, random_state=SEED)
    wb = Workbook()
    ins = wb.active
    ins.title = "Instructions"
    text = ["Word-category audit: 180 words", "",
            "Each row is one word, shown in the line it stands in. The category was set by rules over TLHdig's own "
            "morphological analyses (no language model). Please judge the category of the word in this line.",
            "", "For each row on the 'Audit' sheet:",
            "  1. 'correct? (y/n)': y if the proposed category is right, n if it is wrong.",
            "  2. If n, pick the right one under 'correct category'.",
            "  3. 'note' is optional (e.g. 'adjective', 'place name', 'reading uncertain').",
            "Leave a row empty if you cannot judge it; partial sheets can be scored too.", "",
            "Categories (StBoT 73, Table 4):"] + [f"  {c}: {h}" for c, h in CLASS_HELP] + [
            "", "When done, save the file under the same name in the shared folder. It is read back with",
            "  python review_wordclass.py collect <file>  and scored with  python review_wordclass.py score",
            "which gives the accuracy per confidence level and an estimate for the whole corpus."]
    for i, line in enumerate(text, 1):
        ins.cell(i, 1, line).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    ins.column_dimensions["A"].width = 120
    ws = wb.create_sheet("Audit")
    ws.append(FORM_COLS)
    for r in a.itertuples():
        ws.append([int(r.id), r.txtid, r.lnr, r.line, r.translit, r.wclass, "", "", ""])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EDEBE6")
    for col, w in zip("ABCDEFGHI", (5, 16, 18, 60, 22, 16, 13, 20, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    n = ws.max_row
    yn = DataValidation(type="list", formula1='"y,n"', allow_blank=True)
    cl = DataValidation(type="list", formula1='"' + ",".join(c for c, _ in CLASS_HELP) + '"', allow_blank=True)
    ws.add_data_validation(yn)
    ws.add_data_validation(cl)
    yn.add(f"G2:G{n}")
    cl.add(f"H2:H{n}")
    ws.freeze_panes = "A2"
    FORM.parent.mkdir(exist_ok=True)
    wb.save(FORM)
    print(f"{FORM}: {n - 1} rows")


def collect(path):
    """Copy verdicts from a filled form into the audit csv (matched on the row id)."""
    f = pd.read_excel(path, sheet_name="Audit", dtype=str).fillna("")
    a = pd.read_csv(AUDIT, dtype=str).fillna("")
    for r in f.itertuples(index=False):
        i = int(r[0])
        v, c, note = str(r[6]).strip().lower(), str(r[7]).strip().lower(), str(r[8]).strip()
        if v in ("y", "n"):
            a.loc[i, "verdict"] = v
        if c and c not in VALID:
            raise SystemExit(f"row {i}: unknown category {c!r}")
        a.loc[i, "correct_class"] = c
        if note:
            a.loc[i, "note"] = note
    a.to_csv(AUDIT, index=False)
    print(f"{AUDIT}: {int(a.verdict.isin(['y', 'n']).sum())} of {len(a)} marked")


def wilson(k, n, z=1.96):
    if n == 0:
        return None
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return [round(c - h, 3), round(c + h, 3)]


def score(t):
    a = pd.read_csv(AUDIT, dtype=str).fillna("")
    done = a[a.verdict.str.lower().isin(["y", "n"])]
    share = t.wc_conf.value_counts(normalize=True)
    out = {"reviewed": int(len(done)), "of": int(len(a)), "levels": {}}
    est = 0.0
    for lev, g in done.groupby("wc_conf"):
        k, n = int((g.verdict.str.lower() == "y").sum()), len(g)
        out["levels"][lev] = {"reviewed": n, "right": k, "accuracy": round(k / n, 3), "ci95": wilson(k, n),
                              "share_of_tokens": round(float(share.get(lev, 0)), 3),
                              "errors_by_class": g[g.verdict.str.lower() == "n"].wclass.value_counts().to_dict()}
        est += k / n * share.get(lev, 0)
    if set(out["levels"]) >= {"high", "medium", "low"}:
        out["corpus_accuracy_estimate"] = round(float(est / share[["high", "medium", "low"]].sum()), 3)
    SCORE.parent.mkdir(exist_ok=True)
    SCORE.write_text(json.dumps(out, indent=2, ensure_ascii=False))
    print(json.dumps(out, indent=1, ensure_ascii=False))


def rescore(t):
    """The audit against the current token table: a row is right when its class now is the class
    the philologist accepted (the proposed one if marked y, the corrected one if marked n). Rows
    keep the confidence level they were drawn from, so the estimate weights them as `score` does.
    Rules fitted to these marks make this an in-sample figure: read it as an upper bound."""
    a = pd.read_csv(AUDIT, dtype=str).fillna("")
    a = a[a.verdict.str.lower().isin(["y", "n"])].copy()
    a["accepted"] = np.where(a.verdict.str.lower() == "y", a.wclass, a.correct_class)
    now = t.drop_duplicates(["txtid", "lnr", "translit"])[["txtid", "lnr", "translit", "wclass", "wc_rule"]]
    a = a.merge(now.rename(columns={"wclass": "now", "wc_rule": "rule_now"}), on=["txtid", "lnr", "translit"], how="left")
    # the audit was marked in the earlier scheme, where every noun and adjective was `realia`: score
    # the meaning classes of nouns (and a meaning tie among them) as that one class
    noun = {"realia", "place", "animal", "body", "nature", "abstract", "adjective", "measure"}
    coarse = np.where(a.now.isin(noun) | ((a.now == "ambiguous") & (a.rule_now != "analysis_ambiguous")),
                      "realia", a.now)
    a["right"] = (a.now == a.accepted) | (coarse == a.accepted)
    share = {l: v["share_of_tokens"] for l, v in json.loads(SCORE.read_text())["levels"].items()}  # as drawn
    out = {"reviewed": int(len(a)), "levels": {}, "note": "in-sample: the rules were revised on these marks"}
    est = 0.0
    for lev, g in a.groupby("wc_conf"):
        k, n = int(g.right.sum()), len(g)
        out["levels"][lev] = {"reviewed": n, "right": k, "accuracy": round(k / n, 3), "ci95": wilson(k, n)}
        est += k / n * share.get(lev, 0)
    out["corpus_accuracy_estimate"] = round(float(est / sum(share.get(l, 0) for l in out["levels"])), 3)
    out["still_wrong"] = a[~a.right][["word", "translit", "wclass", "accepted", "now", "rule_now", "note"]].to_dict("records")
    f = HERE / "results" / "wordclass_audit_after_fixes.json"
    f.write_text(json.dumps(out, indent=2, ensure_ascii=False))
    print(json.dumps({k: v for k, v in out.items() if k != "still_wrong"}, indent=1, ensure_ascii=False))
    print(f"{f}: {len(out['still_wrong'])} rows still differ from the philologist's class")
    scores(out, t, noun)


AUDIT_FIXES = 5   # revised on the marks: {M}/{F} names, compound-logogram heads, determinatives
#                   of words without analysis, broken enclitic chains, broken -anzi verbs (qtm/wordclass.py)
FIX_RULES = {"determinative_agent", "determinative_noun", "broken_clitic", "broken_ending"}


def scores(after, t, noun):
    """One file with the audit before and after the rule revisions, for the paper: accuracy by
    confidence level, the corpus-weighted estimate, and the share of words the revisions moved.
    The share counts the words whose class now comes from a revised rule (FIX_RULES); the meaning
    classes of nouns are not fixes, and their share is given apart."""
    before = json.loads(SCORE.read_text())
    lev = lambda d: {k: {x: v[x] for x in ("reviewed", "right", "accuracy", "ci95")} for k, v in d["levels"].items()}
    ch = pd.read_csv(HERE / "results" / "wordclass_rule_changes_audit.csv")
    legible = int((pd.read_csv(t.attrs["path"], usecols=["word"], dtype=str).word.fillna("") != "").sum())
    fixes = ch[ch.rule.isin(FIX_RULES)].tokens.sum()   # the rest moved by meaning class or by evidence
    out = {"sample_size": before["reviewed"], "sample_drawn": before["of"],
           "before_fixes": {"levels": lev(before), "corpus_weighted_accuracy": before["corpus_accuracy_estimate"]},
           "after_fixes": {"levels": lev(after), "corpus_weighted_accuracy": after["corpus_accuracy_estimate"],
                           "in_sample": True, "rows_still_wrong": len(after["still_wrong"])},
           "weights_share_of_tokens": {k: v["share_of_tokens"] for k, v in before["levels"].items()},
           "rule_fixes": AUDIT_FIXES,
           "share_of_words_changed_by_fixes": round(float(fixes / legible), 4),
           "share_of_words_changed_incl_meaning_classes": round(float(ch.tokens.sum() / legible), 4),
           "realia_split_included": True,
           "note": "after-fix accuracy is in-sample (rules revised on these marks) and scored in the audit's "
                   "scheme, where nouns and adjectives were one class (realia); the realia split into meaning "
                   "classes (philologist's decision 2026-10-17) uses labels/gloss_lexicon.tsv, drafted by Claude "
                   "from TLHdig's German glosses and not audited"}
    f = HERE / "results" / "wordclass_audit_scores.json"
    f.write_text(json.dumps(out, indent=2, ensure_ascii=False))
    print(f"{f}: {out['share_of_words_changed_by_fixes']:.1%} of words changed by the fixes, "
          f"{out['share_of_words_changed_incl_meaning_classes']:.1%} with the meaning classes")


def compare(old, new, out=None):
    key = ["txtid", "lnr", "translit", "n"]
    load_ = lambda f, cols: (lambda d: d.assign(n=d.groupby(key[:3], dropna=False).cumcount()))(
        pd.read_csv(f, low_memory=False, usecols=key[:3] + cols, dtype=str))
    a = load_(old, ["word", "wclass"])
    b = load_(new, ["word", "wclass", "wc_rule"])
    b = b.merge(a.rename(columns={"word": "old_word", "wclass": "old_class"}), on=key, how="inner")
    a = b.assign(word=b.old_word, wclass=b.old_class)
    legible = b.word.notna() & (b.word != "")
    ch = legible & (a.wclass.fillna("") != b.wclass.fillna(""))
    d = pd.DataFrame({"from": a.wclass[ch].fillna("(illegible)"), "to": b.wclass[ch], "rule": b.wc_rule[ch],
                      "word": b.word[ch], "translit": b.translit[ch]})
    g = d.groupby(["from", "to", "rule"]).agg(tokens=("word", "size"),
                                              examples=("word", lambda s: ", ".join(s.value_counts().index[:5])))
    g = g.reset_index().sort_values("tokens", ascending=False)
    out = Path(out) if out else HERE / "results" / "wordclass_rule_changes.csv"
    g.to_csv(out, index=False)
    print(f"{out}: {int(ch.sum())} of {int(legible.sum())} legible tokens changed class ({ch.sum() / legible.sum():.1%})")
    print(g.head(20).to_string(index=False))


def examples(t, per=3):
    """Examples per (rule, class), with the raw TLHdig analysis read back from the XML."""
    from lxml import etree
    rows = []
    for (rule, cls), g in t.groupby(["wc_rule", "wclass"]):
        top = g.word.value_counts().index[:per]
        for w in top:
            r = g[g.word == w].iloc[0]
            rows.append({"rule": rule, "class": cls, "confidence": r.wc_conf, "word": w, "translit": r.translit,
                         "txtid": r.txtid, "lnr": r.lnr, "tokens": int((g.word == w).sum())})
    ex = pd.DataFrame(rows)
    # the analysis attributes of the first <w> with that transliteration in that manuscript
    root = HERE / "data" / "TLHbasisONLINE25_1_ZENODO_Beta_03"
    files = {p.stem: p for p in root.rglob("*.xml")}
    ana = []
    for r in ex.itertuples():
        s = ""
        f = files.get(r.txtid)
        if f is not None:
            tree = etree.parse(str(f), etree.XMLParser(recover=True, huge_tree=True))
            for w in tree.iter("w"):
                tr = w.get("trans") or ""
                if tr == r.word or tr.endswith(" " + r.word) or tr.startswith(r.word + " "):
                    s = " | ".join(f"{k}={v}" for k, v in w.attrib.items() if k.startswith("mrp") or k == "lg")
                    break
        ana.append(s[:400])
    ex["analysis"] = ana
    out = HERE / "results" / "wordclass_examples.csv"
    ex.to_csv(out, index=False)
    share = t.wc_conf.value_counts(normalize=True).round(3).to_dict()
    (HERE / "results" / "wordclass_confidence.json").write_text(json.dumps(
        {"scope": "legible words of the texts the experiments use (reference traditions, residue, orphan pools)",
         "legible_tokens": int(len(t)), "confidence_share": share,
         "rule_share": t.wc_rule.value_counts(normalize=True).round(4).to_dict(),
         "class_share": t.wclass.value_counts(normalize=True).round(4).to_dict()}, indent=2, ensure_ascii=False))
    print(f"{out}: {len(ex)} examples; confidence {share}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("step", choices=["sheet", "apply", "audit", "form", "collect", "score", "rescore", "compare", "examples"])
    ap.add_argument("files", nargs="*", help="compare: OLD.csv NEW.csv [OUT.csv]")
    ap.add_argument("--tokens", default=str(HERE / "data" / "tokens.csv"))
    ap.add_argument("--top", type=int, default=2000, help="rows in the review sheet")
    a = ap.parse_args()
    if a.step == "apply":
        return apply()
    if a.step == "compare":
        return compare(*a.files)
    if a.step == "form":
        return form(None)
    if a.step == "collect":
        return collect(a.files[0])
    t = tokens(a.tokens)
    t.attrs["path"] = a.tokens
    sheet(t, a.top) if a.step == "sheet" else {"audit": audit, "score": score, "rescore": rescore, "examples": examples}[a.step](t)


if __name__ == "__main__":
    main()
